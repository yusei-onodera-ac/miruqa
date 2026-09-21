"""仕様書取り込み: txt/md/docx/pdf を SpecItem に分解する。

流れ:
  1. 形式ごとにテキストを取り出し、見出し(セクション)で分割する(コード側の処理。LLMに頼らない)。
     - md: `#`〜`######` 見出し
     - txt: 「1.」「2.1」のような番号見出しをコードで検出
     - docx: 段落スタイル(Heading/Title)
     - pdf: ページ単位でテキストを取り出し、txtと同じ番号見出し検出をページ内で行う
       (スキャン画像PDFはテキストが取れないため対象外。warningsに残す)
  2. 各セクションの本文をLLMに送る前にPIIマスクする(絶対条件7)。
  3. セクションごとにLLM(構造化出力・tool_choiceで強制)を呼び、規範的な記述(SpecItem)を抜き出す。
     section/pageはコード側で既に分かっているため、LLMにはtext(内容)とkind(分類)だけ出させる
     (仕様項目IDとの対応=judgement根拠を、なるべくコード側で担保する)。
  4. 章数が多い仕様書は SPEC_MAX_CHUNKS で打ち切り、warnings に残す(コスト上限)。

失敗しても落ちない: 個々のセクション抽出が失敗(LLM呼び出し失敗)しても、そのセクションを
スキップしてwarningsに記録し、処理を続ける(FR-30)。
"""

import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from . import config, ledger, llm, masking

SUPPORTED_TYPES = {".txt": "txt", ".md": "md", ".docx": "docx", ".pdf": "pdf"}
SPEC_ITEM_KINDS = ("rule", "flow", "ui", "validation", "constraint", "other")

EXTRACT_TOOL = {
    "type": "function",
    "function": {
        "name": "extract_spec_items",
        "description": (
            "渡されたセクションの文章から、仕様として守るべき規範的な記述"
            "(〜しなければならない、〜できない、〜とする、等)を項目として抜き出す。"
            "単なる説明文で規範的な記述がなければ、空配列を返す。1セクションに複数の規範的な"
            "記述があれば、複数の項目に分けて返す。"
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "items": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "text": {"type": "string", "description": "仕様項目の内容を1文で(原文の言い回しを保ちつつ簡潔に)"},
                            "kind": {"type": "string", "enum": list(SPEC_ITEM_KINDS)},
                        },
                        "required": ["text", "kind"],
                    },
                }
            },
            "required": ["items"],
        },
    },
}

EXTRACT_SYSTEM_PROMPT = """\
あなたは仕様書からテスト項目書を作るQA担当です。渡されたセクションの見出しと本文を読み、
「実装がこの通りに動くべき」という規範的な記述だけを extract_spec_items で抜き出してください。
背景説明・概要・単なる見出しの繰り返しなど、規範的な記述を含まないセクションは items を空配列にしてください。
本文に書かれていないことを推測で付け足さないでください。"""


_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(\S.{0,38})$")


def _flush(sections, title, body_lines):
    body = "\n".join(body_lines).strip()
    if title is not None or body:
        sections.append((title, body))


def _split_markdown(text):
    sections = []
    title, body = None, []
    for line in text.splitlines():
        m = _MD_HEADING_RE.match(line)
        if m:
            _flush(sections, title, body)
            title, body = m.group(2).strip(), []
        else:
            body.append(line)
    _flush(sections, title, body)
    return sections


def _split_plain_text(text):
    sections = []
    title, body = None, []
    for line in text.splitlines():
        stripped = line.strip()
        m = _NUMBERED_HEADING_RE.match(stripped)
        if m and len(stripped) <= 40:
            _flush(sections, title, body)
            title, body = stripped, []
        else:
            body.append(line)
    _flush(sections, title, body)
    return sections


def _read_docx_sections(path):
    from docx import Document

    doc = Document(str(path))
    sections = []
    title, body = None, []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style else "") or ""
        if style.startswith("Heading") or style == "Title":
            _flush(sections, title, body)
            title, body = text, []
        else:
            body.append(text)
    _flush(sections, title, body)
    warnings = []
    if doc.tables:
        warnings.append(f"表を{len(doc.tables)}個検出しましたが、現時点では表の内容は読み取っていません。")
    return [(t, b, None) for t, b in sections], warnings


def _read_pdf_sections(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections = []
    warnings = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if not text.strip():
            warnings.append(f"{i}ページ目からテキストを抽出できませんでした(スキャン画像の可能性があります。OCRは非対応です)。")
            continue
        for title, body in _split_plain_text(text):
            if body.strip() or title:
                sections.append((title, body, i))
    return sections, warnings


def extract_sections(path):
    """(sections: list[(title, body, page)], warnings, doc_type) を返す。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_TYPES:
        raise ValueError(f"未対応の形式です: {ext}(対応形式: {', '.join(SUPPORTED_TYPES)})")
    doc_type = SUPPORTED_TYPES[ext]

    if doc_type in ("txt", "md"):
        raw = path.read_text(encoding="utf-8", errors="replace")
        splitter = _split_markdown if doc_type == "md" else _split_plain_text
        sections = [(t, b, None) for t, b in splitter(raw)]
        return sections, [], doc_type
    if doc_type == "docx":
        sections, warnings = _read_docx_sections(path)
        return sections, warnings, doc_type
    sections, warnings = _read_pdf_sections(path)
    return sections, warnings, doc_type


def build_spec(file_path, filename=None, client=None):
    """仕様書ファイル1件を SpecItem[] に分解する。戻り値: (spec: dict, llm_calls: list[dict])。"""
    path = Path(file_path)
    sections, warnings, doc_type = extract_sections(path)

    if len(sections) > config.SPEC_MAX_CHUNKS:
        warnings.append(f"仕様書のセクション数({len(sections)})が上限({config.SPEC_MAX_CHUNKS})を超えたため、先頭{config.SPEC_MAX_CHUNKS}件のみ処理しました。")
        sections = sections[: config.SPEC_MAX_CHUNKS]

    spec_id = f"s-{uuid.uuid4().hex[:8]}"  # 台帳(runs/ledger.jsonl)への紐付け・予算スコープのため先に確定させる
    client = client or llm.new_client()
    items = []
    llm_calls = []
    consecutive_degraded = 0

    for title, body, page in sections:
        text = (body or "").strip()
        if not text:
            continue
        if consecutive_degraded >= 2:
            # 2セクション連続でLLM呼び出しが復旧できない場合、系統的な障害(またはコスト上限到達)と
            # みなして残りセクションのリトライを打ち切る(全セクション分の待ち時間を防ぐ。FR-30)。
            last_error = (llm_calls[-1].get("error") if llm_calls else "") or ""
            if last_error.startswith("budget_exceeded"):
                warnings.append(f"コスト上限に達したため、残りのセクションの抽出を打ち切りました({last_error})。")
            else:
                warnings.append("LLM呼び出しが連続して失敗したため、残りのセクションの抽出を打ち切りました。")
            break
        masked_text, _ = masking.mask_text(text)
        messages = [
            {"role": "system", "content": EXTRACT_SYSTEM_PROMPT},
            {"role": "user", "content": f"セクション見出し: {title or '(なし)'}\n本文:\n{masked_text}"},
        ]
        message, llm_call, degraded = llm.call_llm(
            client,
            messages=messages,
            purpose="spec-extract",
            tools=[EXTRACT_TOOL],
            tool_choice={"type": "function", "function": {"name": "extract_spec_items"}},
            context={"specId": spec_id},
        )
        llm_calls.append(llm_call)

        if degraded or message is None or not message.tool_calls:
            consecutive_degraded += 1
            warnings.append(f"セクション「{title or '(見出しなし)'}」の抽出に失敗したためスキップしました({llm_call.get('error') or '応答がツール呼び出しではありませんでした'})。")
            continue
        consecutive_degraded = 0

        args = llm.safe_json_loads(message.tool_calls[0].function.arguments, default={})
        for it in args.get("items", []) or []:
            text_val = (it.get("text") or "").strip()
            if not text_val:
                continue
            kind = it.get("kind") if it.get("kind") in SPEC_ITEM_KINDS else "other"
            items.append(
                {
                    "id": f"SPEC-{len(items) + 1:03d}",
                    "text": text_val,
                    "section": title,
                    "page": page,
                    "kind": kind,
                }
            )

    spec = {
        "specId": spec_id,
        "filename": filename or path.name,
        "type": doc_type,
        "uploadedAt": datetime.now(timezone.utc).isoformat(),
        "items": items,
        "warnings": warnings,
        "budgetStatus": ledger.budget_status_from_calls(llm_calls),
    }
    return spec, llm_calls


def save_spec(spec, source_path=None):
    """runs/specs/<specId>/spec.json に保存する(サーバー・CLI共通の永続化)。"""
    spec_dir = config.SPECS_DIR / spec["specId"]
    spec_dir.mkdir(parents=True, exist_ok=True)
    if source_path is not None:
        source_path = Path(source_path)
        dest = spec_dir / source_path.name
        if source_path.resolve() != dest.resolve():
            dest.write_bytes(source_path.read_bytes())
    (spec_dir / "spec.json").write_text(json.dumps(spec, ensure_ascii=False, indent=1), encoding="utf-8")
    return spec_dir


def load_spec(spec_id):
    path = config.SPECS_DIR / spec_id / "spec.json"
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def load_specs(spec_ids):
    """複数のSpecIdからSpecItem辞書(id -> item)をまとめて作る(TestCase生成で使う)。存在しないIDは無視する。"""
    items_by_id = {}
    specs = []
    for spec_id in spec_ids or []:
        spec = load_spec(spec_id)
        if not spec:
            continue
        specs.append(spec)
        for item in spec.get("items", []):
            items_by_id[item["id"]] = item
    return specs, items_by_id


def main(argv=None):
    import argparse

    parser = argparse.ArgumentParser(description="仕様書をSpecItemに分解する(単体テスト用CLI)")
    parser.add_argument("path", help="仕様書ファイル(txt/md/docx/pdf)")
    parser.add_argument("--save", action="store_true", help="runs/specs/ に保存する")
    args = parser.parse_args(argv)

    spec, llm_calls = build_spec(args.path)
    print(json.dumps(spec, ensure_ascii=False, indent=1))
    print(f"\n項目数: {len(spec['items'])} / セクション数(LLM呼び出し回数): {len(llm_calls)}", flush=True)
    if args.save:
        spec_dir = save_spec(spec, source_path=args.path)
        print(f"保存先: {spec_dir}")


if __name__ == "__main__":
    main()
